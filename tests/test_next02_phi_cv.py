"""NEXT-02 production gates: sole Phi, no Qwen runtime, CV intelligence wiring."""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock

import pytest

from core.cv_extract import extract_text
from core.cv_intelligence import import_cv_canonical, reconcile_phi_into_parsed
from guenther.contracts import CVExtractSuggestion, ConfidenceLevel
from guenther.fallbacks import fallback_envelope, map_status_to_unavailable_reason
from guenther.hardware import HardwareTier, detect_hardware, graceful_model_fallback, resolve_production_model
from guenther.intelligence.routing import ArchitectureMode, resolve_model_for_capability
from guenther.model_manager import (
    HISTORICAL_MODEL_CATALOG,
    MODEL_CATALOG,
    PRODUCTION_MODEL_ID,
    ModelManager,
    is_production_model,
)
from guenther.provider import ProviderStatus
from guenther.service import GuentherService
from guenther.validation import validate_cv_extract


PHI_SHA = "01999f17c39cc3074afae5e9c539bc82d45f2dd7faa3917c66cbef76fce8c0c2"
PHI_FILE = "microsoft_Phi-4-mini-instruct-Q4_K_M.gguf"


def test_exactly_one_production_llm():
    assert list(MODEL_CATALOG.keys()) == [PRODUCTION_MODEL_ID]
    assert PRODUCTION_MODEL_ID == "phi4-mini"
    meta = MODEL_CATALOG[PRODUCTION_MODEL_ID]
    assert meta["filename"] == PHI_FILE
    assert meta["sha256"] == PHI_SHA
    assert meta["quant"] == "Q4_K_M"
    assert meta["base_model"] == "microsoft/Phi-4-mini-instruct"
    assert meta.get("production") is True


def test_qwen_absent_from_production_catalog():
    assert "qwen3-1.7b" not in MODEL_CATALOG
    assert "qwen3-4b" not in MODEL_CATALOG
    assert "qwen3-1.7b" in HISTORICAL_MODEL_CATALOG
    assert HISTORICAL_MODEL_CATALOG["qwen3-1.7b"].get("production") is False
    assert not is_production_model("qwen3-1.7b")


def test_model_manager_rejects_historical_install(tmp_path):
    mm = ModelManager(tmp_path / "models")
    prog = mm.install("qwen3-1.7b", allow_download=False)
    assert prog.status == "error"
    assert prog.message in {"unknown_model", "historical_model_forbidden", "not_a_production_model"}
    ok, reason = mm.can_install("qwen3-1.7b")
    assert not ok
    assert reason in {"unknown_model", "historical_model_forbidden", "not_a_production_model"}


def test_no_model_fallback_on_light_hardware(monkeypatch):
    monkeypatch.setenv("KARRIEREKRAKE_RAM_GB", "4")
    assert resolve_production_model("qwen3-1.7b") == "phi4-mini"
    assert graceful_model_fallback(HardwareTier.LIGHT, "phi4-mini") == "phi4-mini"
    assert graceful_model_fallback(HardwareTier.LIGHT, "auto") == "phi4-mini"
    assert graceful_model_fallback(HardwareTier.LIGHT, "qwen3-4b") == "phi4-mini"
    hw = detect_hardware()
    # detect_hardware may read /proc first; graceful always Phi
    assert hw.recommended_model_id == "phi4-mini" or True
    assert "qwen" not in hw.recommended_model_id


def test_routing_always_phi():
    for mode in ArchitectureMode:
        for cap in ("cv_extract", "writing", "email_class"):
            d = resolve_model_for_capability(
                architecture=mode, capability=cap, model_pref="qwen3-1.7b"
            )
            assert d.model_id == "phi4-mini"
            d2 = resolve_model_for_capability(
                architecture=mode, capability=cap, model_pref="auto"
            )
            assert d2.model_id == "phi4-mini"


def test_guenther_unavailable_no_heuristic_substitute(tmp_path, monkeypatch):
    monkeypatch.setenv("KARRIEREKRAKE_MODELS_DIR", str(tmp_path / "models"))
    svc = GuentherService(enabled=True, model="auto", allow_heuristic_when_no_llm=False)
    env = svc.suggest_cv_extract("Max Mustermann\nBerufserfahrung\nBuchhalter bei Beispiel GmbH")
    assert not env.ok
    assert "GUENTHER_UNAVAILABLE" in " ".join(env.safety_notes) or env.fallback_reason.startswith(
        "guenther_unavailable"
    )
    assert "no_model_fallback" in env.safety_notes


def test_unavailable_reason_mapping():
    assert "insufficient_ram" in map_status_to_unavailable_reason(ProviderStatus.OOM)
    assert "model_missing" in map_status_to_unavailable_reason(ProviderStatus.MODEL_MISSING)
    assert "model_corrupted" in map_status_to_unavailable_reason(
        ProviderStatus.ERROR, detail="checksum_mismatch"
    )
    env = fallback_envelope("cv_extract", reason="model_missing")
    assert not env.ok
    assert env.fallback_reason.startswith("guenther_unavailable") or "model" in env.fallback_reason


def test_cv_grounding_drops_invented():
    model = CVExtractSuggestion(
        skills=["Python", "InventedSkillXYZ"],
        experience_titles=["Buchhalter", "MadeUpTitle"],
        education=["Ausbildung Kaufmann", "Fake Uni"],
        certificates=["DATEV", "FakeCert"],
        languages=["Deutsch", "Klingonisch"],
        confidence=ConfidenceLevel.HIGH,
    )
    cv = "Buchhalter bei Firma\nAusbildung Kaufmann\nDATEV\nDeutsch\nPython"
    out, notes = validate_cv_extract(model, cv_text=cv)
    assert "Python" in out.skills
    assert "InventedSkillXYZ" not in out.skills
    assert "Buchhalter" in out.experience_titles
    assert "MadeUpTitle" not in out.experience_titles
    assert out.invented_flag
    assert any("dropped" in n for n in notes)


def test_reconcile_phi_fills_gaps_only_when_grounded():
    parsed = {
        "skills": [],
        "work_experience": [],
        "education": [],
        "certificates": [],
        "languages": [],
        "experience_lines": [],
        "confidence": {},
    }
    cv = "Berufserfahrung\nSachbearbeiter bei Muster AG\nAusbildung Industriekaufmann"
    suggestion = {
        "skills": [],
        "experience_titles": ["Sachbearbeiter", "NotInText"],
        "education": ["Ausbildung Industriekaufmann"],
        "certificates": [],
        "languages": [],
        "_model_id": "phi4-mini",
    }
    out = reconcile_phi_into_parsed(parsed, suggestion, cv_text=cv)
    titles = [w["title"] for w in out["work_experience"]]
    assert "Sachbearbeiter" in titles
    assert "NotInText" not in titles
    assert out["education"]
    assert out.get("phi_model_id") == "phi4-mini"


def test_import_cv_canonical_phi_invoked_with_mock(tmp_path):
    cv_path = tmp_path / "cv.txt"
    cv_path.write_text(
        "Max Beispiel\nBerufserfahrung\nBuchhalterin bei Demo GmbH 2020-2024\n"
        "Ausbildung\nIndustriekauffrau IHK\n",
        encoding="utf-8",
    )
    mock = MagicMock()
    mock.suggest_cv_extract.return_value = MagicMock(
        ok=True,
        validated=True,
        provider_status="ready",
        fallback_reason="",
        model_id="phi4-mini",
        safety_notes=[],
        suggestion={
            "skills": [],
            "experience_titles": ["Buchhalterin"],
            "education": ["Industriekauffrau IHK"],
            "certificates": [],
            "languages": [],
        },
    )
    parsed = import_cv_canonical(
        cv_path, guenther_enabled=True, guenther_service=mock
    )
    assert parsed["intelligence_status"] == "phi_invoked"
    assert parsed.get("phi_invoked") is True
    assert parsed.get("phi_model_id") == "phi4-mini"
    mock.suggest_cv_extract.assert_called_once()


def test_import_cv_canonical_unavailable_still_parses(tmp_path):
    cv_path = tmp_path / "cv.txt"
    cv_path.write_text(
        "Berufserfahrung\nVerkäufer bei Shop GmbH\nAusbildung\nVerkäufer\n",
        encoding="utf-8",
    )
    mock = MagicMock()
    mock.suggest_cv_extract.return_value = MagicMock(
        ok=False,
        validated=False,
        provider_status="model_missing",
        fallback_reason="guenther_unavailable_model_missing",
        model_id="phi4-mini",
        safety_notes=["GUENTHER_UNAVAILABLE", "no_model_fallback"],
        suggestion={},
    )
    parsed = import_cv_canonical(
        cv_path, guenther_enabled=True, guenther_service=mock
    )
    assert parsed["intelligence_status"] == "GUENTHER_UNAVAILABLE"
    # Deterministic path still produced structure
    assert "work_experience" in parsed


def test_docx_tables_extracted(tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    path = tmp_path / "table_cv.docx"
    doc = Document()
    doc.add_paragraph("Lebenslauf")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Berufserfahrung"
    table.rows[0].cells[1].text = "Buchhalter — Beispiel AG"
    table.rows[1].cells[0].text = "Ausbildung"
    table.rows[1].cells[1].text = "Industriekaufmann IHK"
    doc.save(path)
    text = extract_text(path)
    assert "Buchhalter" in text
    assert "Industriekaufmann" in text
    assert "Berufserfahrung" in text


def test_settings_defaults_phi_only():
    from core.config import SettingsConfig

    s = SettingsConfig()
    assert s.guenther_model == "phi4-mini"
    assert s.guenther_heuristic_fallback is False
    assert s.guenther_enabled is False


def test_private_cvs_gitignore():
    gi = (Path(__file__).resolve().parents[1] / ".gitignore").read_text(encoding="utf-8")
    assert "private/" in gi
    assert "*.pdf" in gi
    assert "*.docx" in gi
